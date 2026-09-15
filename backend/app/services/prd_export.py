import io
import re
from datetime import date

from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    HRFlowable,
)

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


PRD_SECTIONS = [
    ("Product Overview", "product_overview"),
    ("Problem Statement", "prd_problem_statement"),
    ("Goals and Objectives", "goals_and_objectives"),
    ("Target Users", "prd_target_users"),
    ("Functional Requirements", "functional_requirements"),
    ("User Workflows", "user_workflows"),
    ("Business Rules", "business_rules"),
    ("Non-Functional Requirements", "non_functional_requirements"),
    ("Assumptions and Dependencies", "assumptions_and_dependencies"),
    ("Acceptance Criteria", "acceptance_criteria"),
]

ACCENT_HEX = "#5B47E0"
TEXT_HEX = "#1F2430"
MUTED_HEX = "#6B7280"


# =========================================================
# TEXT SANITIZATION
# =========================================================
#
# Base PDF/Word fonts don't have glyphs for arrows, fancy
# bullets, checkmarks etc. Unsupported characters render as
# black boxes. Map the common ones LLMs tend to produce to
# safe equivalents.
# =========================================================

CHAR_REPLACEMENTS = {
    "\u2192": "->",   # →
    "\u2190": "<-",   # ←
    "\u2194": "<->",  # ↔
    "\u21d2": "=>",   # ⇒
    "\u2713": "[x]",  # ✓
    "\u2714": "[x]",  # ✔
    "\u2716": "[x]",  # ✖
    "\u2022": "-",    # •
    "\u25cf": "-",    # ●
    "\u25aa": "-",    # ▪
    "\u2023": "-",    # ‣
    "\u2605": "*",    # ★
    "\u2606": "*",    # ☆
    "\u2026": "...",  # …
    "\u2018": "'",
    "\u2019": "'",
    "\u201c": '"',
    "\u201d": '"',
}


def sanitize_text(text: str) -> str:

    if not text:
        return ""

    for bad, good in CHAR_REPLACEMENTS.items():
        text = text.replace(bad, good)

    # Drop anything else outside safe Latin-1 range
    # (keeps standard accented characters, strips emoji/
    # exotic symbols/box-drawing characters etc.)
    text = "".join(
        ch if ord(ch) <= 0x00FF or ch in "\n\t" else ""
        for ch in text
    )

    return text


# =========================================================
# CONTENT PARSING
# =========================================================
#
# Splits LLM output into typed blocks so bullets/numbered
# lists render as real lists, not raw dashes.
# =========================================================

BULLET_RE = re.compile(r"^\s*[-*]\s+(.*)")
NUMBER_RE = re.compile(r"^\s*(\d+)[\.\)]\s+(.*)")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
INLINE_NUMBER_SPLIT_RE = re.compile(r"(?<=[.\)])\s+(?=\d+[\.\)]\s)")
INLINE_BULLET_SPLIT_RE = re.compile(r"(?<=[.\)])\s+(?=[-*]\s+\S)")


def split_inline_lists(text: str) -> str:

    # Break "...cart. 4. Order placement..." into separate lines
    # before "4." when it follows the end of a sentence.
    text = INLINE_NUMBER_SPLIT_RE.sub("\n", text)

    # Same idea for inline "- " bullets glued onto one line.
    text = INLINE_BULLET_SPLIT_RE.sub("\n", text)

    return text

def parse_blocks(raw_text: str):

    text = sanitize_text(raw_text or "")

    text = split_inline_lists(text)

    blocks = []

    for line in text.split("\n"):

        line = line.strip()

        if not line:
            continue

        bullet_match = BULLET_RE.match(line)
        number_match = NUMBER_RE.match(line)

        if bullet_match:

            blocks.append(("bullet", bullet_match.group(1).strip()))

        elif number_match:

            blocks.append(("number", number_match.group(2).strip()))

        else:

            blocks.append(("para", line))

    if not blocks:
        blocks.append(("para", "Not available."))

    return blocks

    text = sanitize_text(raw_text or "")

    blocks = []

    for line in text.split("\n"):

        line = line.strip()

        if not line:
            continue

        bullet_match = BULLET_RE.match(line)
        number_match = NUMBER_RE.match(line)

        if bullet_match:

            blocks.append(("bullet", bullet_match.group(1).strip()))

        elif number_match:

            blocks.append(("number", number_match.group(2).strip()))

        else:

            blocks.append(("para", line))

    if not blocks:
        blocks.append(("para", "Not available."))

    return blocks


def markdown_bold_to_reportlab(text: str) -> str:

    text = (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )

    return BOLD_RE.sub(r"<b>\1</b>", text)


# =========================================================
# PDF
# =========================================================

def generate_prd_pdf(project):

    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=LETTER,
        topMargin=0.9 * inch,
        bottomMargin=0.9 * inch,
        leftMargin=0.85 * inch,
        rightMargin=0.85 * inch,
        title=f"{project.project_name} - PRD",
    )

    styles = getSampleStyleSheet()

    eyebrow_style = ParagraphStyle(
        "Eyebrow",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        textColor=HexColor(ACCENT_HEX),
        spaceAfter=6,
    )

    title_style = ParagraphStyle(
        "PRDTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=26,
        leading=30,
        textColor=HexColor(TEXT_HEX),
        alignment=TA_LEFT,
        spaceAfter=4,
    )

    meta_style = ParagraphStyle(
        "PRDMeta",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        textColor=HexColor(MUTED_HEX),
        spaceAfter=20,
    )

    heading_style = ParagraphStyle(
        "PRDHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13.5,
        textColor=HexColor(ACCENT_HEX),
        spaceBefore=18,
        spaceAfter=8,
    )

    body_style = ParagraphStyle(
        "PRDBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=16,
        textColor=HexColor(TEXT_HEX),
        spaceAfter=6,
        alignment=TA_LEFT,
    )

    bullet_style = ParagraphStyle(
        "PRDBullet",
        parent=body_style,
        leftIndent=16,
        bulletIndent=2,
        spaceAfter=4,
    )

    story = []

    story.append(Paragraph("PRODUCT REQUIREMENTS DOCUMENT", eyebrow_style))

    story.append(
        Paragraph(
            sanitize_text(project.project_name) or "Untitled Project",
            title_style,
        )
    )

    generated_on = date.today().strftime("%B %d, %Y")

    story.append(
        Paragraph(
            f"Generated on {generated_on} &nbsp;&middot;&nbsp; "
            f"Category: {sanitize_text(project.category) or 'Not specified'}",
            meta_style,
        )
    )

    story.append(
        HRFlowable(
            width="100%",
            thickness=1.2,
            color=HexColor(ACCENT_HEX),
            spaceAfter=10,
        )
    )

    for heading, attr in PRD_SECTIONS:

        raw_value = getattr(project, attr, None)

        story.append(Paragraph(heading, heading_style))

        blocks = parse_blocks(raw_value)

        number_counter = 0

        for block_type, block_text in blocks:

            formatted = markdown_bold_to_reportlab(block_text)

            if block_type == "bullet":

                story.append(
                    Paragraph(
                        formatted,
                        bullet_style,
                        bulletText="-",
                    )
                )

            elif block_type == "number":

                number_counter += 1

                story.append(
                    Paragraph(
                        formatted,
                        bullet_style,
                        bulletText=f"{number_counter}.",
                    )
                )

            else:

                number_counter = 0

                story.append(Paragraph(formatted, body_style))

    def _footer(canvas, doc_):

        canvas.saveState()

        canvas.setFont("Helvetica", 8.5)
        canvas.setFillColor(HexColor(MUTED_HEX))

        canvas.drawString(
            0.85 * inch,
            0.55 * inch,
            sanitize_text(project.project_name) or "Untitled Project",
        )

        canvas.drawRightString(
            LETTER[0] - 0.85 * inch,
            0.55 * inch,
            f"Page {doc_.page}",
        )

        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)

    buffer.seek(0)

    return buffer


# =========================================================
# DOCX
# =========================================================

def _set_font(run, name="Calibri", size=11, color=None, bold=False):

    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold

    if color:
        run.font.color.rgb = RGBColor.from_string(color)

    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))

    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)

    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def _add_runs_with_bold(paragraph, text, size=11, color=None):

    parts = BOLD_RE.split(text)

    # split() on a capturing group alternates
    # [normal, bold, normal, bold, ...]

    for index, part in enumerate(parts):

        if not part:
            continue

        run = paragraph.add_run(part)

        is_bold = index % 2 == 1

        _set_font(run, size=size, color=color, bold=is_bold)


def generate_prd_docx(project):

    document = Document()

    section = document.sections[0]
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)

    accent = ACCENT_HEX.lstrip("#")
    text_color = TEXT_HEX.lstrip("#")
    muted = MUTED_HEX.lstrip("#")

    project_name = sanitize_text(project.project_name) or "Untitled Project"

    # -----------------------------------------------------
    # EYEBROW
    # -----------------------------------------------------

    eyebrow = document.add_paragraph()
    run = eyebrow.add_run("PRODUCT REQUIREMENTS DOCUMENT")
    _set_font(run, size=10, color=accent, bold=True)

    # -----------------------------------------------------
    # TITLE
    # -----------------------------------------------------

    title_p = document.add_paragraph()
    title_run = title_p.add_run(project_name)
    _set_font(title_run, name="Calibri", size=26, color=text_color, bold=True)
    title_p.space_after = Pt(4)

    # -----------------------------------------------------
    # META LINE
    # -----------------------------------------------------

    generated_on = date.today().strftime("%B %d, %Y")

    meta_p = document.add_paragraph()
    meta_run = meta_p.add_run(
        f"Generated on {generated_on}  |  "
        f"Category: {sanitize_text(project.category) or 'Not specified'}"
    )
    _set_font(meta_run, size=9.5, color=muted)
    meta_p.space_after = Pt(14)

    # -----------------------------------------------------
    # DIVIDER
    # -----------------------------------------------------

    divider = document.add_paragraph()
    divider_run = divider.add_run("-" * 60)
    _set_font(divider_run, size=8, color=accent)
    divider.space_after = Pt(10)

    # -----------------------------------------------------
    # SECTIONS
    # -----------------------------------------------------

    for heading, attr in PRD_SECTIONS:

        raw_value = getattr(project, attr, None)

        heading_p = document.add_paragraph()
        heading_p.space_before = Pt(16)
        heading_p.space_after = Pt(6)

        heading_run = heading_p.add_run(heading)
        _set_font(heading_run, size=13.5, color=accent, bold=True)

        blocks = parse_blocks(raw_value)

        for block_type, block_text in blocks:

            if block_type == "bullet":

                body_p = document.add_paragraph(style="List Bullet")

            elif block_type == "number":

                body_p = document.add_paragraph(style="List Number")

            else:

                body_p = document.add_paragraph()

            body_p.paragraph_format.line_spacing = 1.3
            body_p.space_after = Pt(6)

            _add_runs_with_bold(
                body_p,
                block_text,
                size=11,
                color=text_color,
            )

    # -----------------------------------------------------
    # FOOTER
    # -----------------------------------------------------

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer_run = footer_p.add_run(project_name)
    _set_font(footer_run, size=8.5, color=muted)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)

    return buffer